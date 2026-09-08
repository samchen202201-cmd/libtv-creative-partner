#!/usr/bin/env python3
"""Build a source-derived, landscape production guide from a small Markdown subset."""
from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path
import re
import math
import sys
import unicodedata

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DEFAULT_TEMPLATE = Path(__file__).resolve().parent.parent / 'assets' / 'production-guide.docx'
HEADER_FILL, ROW_FILL, BORDER_COLOR = '7C2D12', 'FFFBEB', 'D0D5DD'


def element(tag, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    return node


def inline(paragraph, text):
    """Support bold spans without interpreting or executing embedded content."""
    for idx, part in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        if part:
            run = paragraph.add_run(part)
            if idx % 2:
                run.bold = True


def cells(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|') and not line.endswith(r'\|'):
        line = line[:-1]
    return [part.replace(r'\|', '|').strip() for part in re.split(r'(?<!\\)\|', line)]


def table_divider(line):
    parts = cells(line)
    return len(parts) > 1 and all(re.fullmatch(r':?-{3,}:?', part) for part in parts)


def parse_markdown(source):
    if not source.strip():
        raise ValueError('Input Markdown is empty.')
    lines, blocks, index = source.splitlines(), [], 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        width_hint = re.fullmatch(r'<!--\s*widths:\s*(.*?)\s*-->', line)
        if width_hint:
            try:
                widths = [float(value.strip()) for value in width_hint[1].split(',')]
            except ValueError as exc:
                raise ValueError(f'Invalid column widths at line {index + 1}.') from exc
            if not widths or any(not math.isfinite(w) or w <= 0 for w in widths):
                raise ValueError(f'Column widths must be finite positive centimetres at line {index + 1}.')
            next_index = index + 1
            while next_index < len(lines) and not lines[next_index].strip():
                next_index += 1
            if next_index + 1 >= len(lines) or '|' not in lines[next_index] or not table_divider(lines[next_index + 1]):
                raise ValueError(f'Column-width comment must immediately precede a table at line {index + 1}.')
            blocks.append(('widths', widths))
            index += 1
            continue
        if re.fullmatch(r'\*\*[^*\n]+\*\*', line):
            blocks.append(('label', line))
            index += 1
            continue
        if line.startswith('```'):
            fence = re.fullmatch(r'(`{3,})([^`]*)', line)
            if fence is None:
                raise ValueError(f'Invalid code fence at line {index + 1}.')
            delimiter, label, start = fence[1], fence[2].strip(), index + 1
            closing = re.compile(r'`{' + str(len(delimiter)) + r',}')
            index += 1
            buffer = []
            while index < len(lines) and not closing.fullmatch(lines[index].strip()):
                buffer.append(lines[index])
                index += 1
            if index == len(lines):
                raise ValueError(f'Unclosed code fence at line {start}.')
            blocks.append(('prompt' if label == 'prompt' else 'literal', '\n'.join(buffer)))
            index += 1
            continue
        if line == '<!-- pagebreak -->':
            blocks.append(('pagebreak', ''))
            index += 1
            continue
        heading = re.fullmatch(r'(#{1,3})\s+(.+)', line)
        if heading:
            blocks.append(('heading' + str(len(heading[1])), heading[2]))
            index += 1
            continue
        if index + 1 < len(lines) and '|' in line and table_divider(lines[index + 1]):
            rows = [cells(line)]
            index += 2
            while index < len(lines) and lines[index].strip() and '|' in lines[index]:
                row = cells(lines[index])
                if len(row) != len(rows[0]):
                    raise ValueError(f'Table column count differs at line {index + 1}.')
                rows.append(row)
                index += 1
            blocks.append(('table', rows))
            continue
        if line.startswith('- '):
            blocks.append(('bullet', line[2:]))
            index += 1
            continue
        if line.startswith('>'):
            blocks.append(('note', line[1:].lstrip()))
            index += 1
            continue
        buffer = [line]
        index += 1
        while index < len(lines) and lines[index].strip():
            following = lines[index].strip()
            if following.startswith(('#', '```', '- ', '>', '<!-- pagebreak -->', '<!-- widths:')) or re.fullmatch(r'\*\*[^*\n]+\*\*', following):
                break
            if index + 1 < len(lines) and '|' in following and table_divider(lines[index + 1]):
                break
            buffer.append(following)
            index += 1
        blocks.append(('paragraph', '\n'.join(buffer)))
    if not any(kind != 'pagebreak' and bool(value) for kind, value in blocks):
        raise ValueError('Input Markdown has no visible content.')
    return blocks


def display_length(text):
    text = re.sub(r'\*\*|<br\s*/?>', '', text)
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in text)


def column_widths(rows, width_cm):
    count = len(rows[0])
    measures = []
    for col in range(count):
        lengths = sorted(display_length(row[col]) for row in rows)
        typical = lengths[min(len(lengths) - 1, int(len(lengths) * .8))]
        measures.append(max(display_length(rows[0][col]), typical))
    compact = {i for i, value in enumerate(measures) if value <= 18}
    widths = [0.0] * count
    for col in compact:
        widths[col] = max(1.35, min(3.3, measures[col] * .12 + .55))
    broad = [i for i in range(count) if i not in compact]
    if not broad:
        # Compact schema: do not distribute every field evenly; keep the first narrow.
        weights = [max(1, m) ** .6 for m in measures]
        if count > 1:
            weights[0] *= .7
        widths = [width_cm * w / sum(weights) for w in weights]
    else:
        # Leave a floor for prose columns before distributing spare width.
        remaining = width_cm - sum(widths)
        if remaining < 1.8 * len(broad):
            scale = (width_cm - 1.8 * len(broad)) / max(sum(widths), .01)
            widths = [w * max(.35, scale) for w in widths]
            remaining = width_cm - sum(widths)
        weights = [min(measures[i], 450) ** .55 for i in broad]
        for col, weight in zip(broad, weights):
            widths[col] = remaining * weight / sum(weights)
    # Keep narrative first columns useful while respecting the compact-key convention.
    if count > 1 and widths[0] > width_cm * .20:
        excess = widths[0] - width_cm * .20
        widths[0] -= excess
        recipients = [i for i in range(1, count) if i in broad] or list(range(1, count))
        total = sum(widths[i] for i in recipients)
        for col in recipients:
            widths[col] += excess * widths[col] / total
    return widths


def add_table(doc, rows, explicit_widths=None):
    section = doc.sections[-1]
    width_cm = (section.page_width - section.left_margin - section.right_margin) / 360000
    if explicit_widths is None:
        widths = column_widths(rows, width_cm)
    else:
        if len(explicit_widths) != len(rows[0]):
            raise ValueError('Column-width count must equal the table column count.')
        # Respect cm values; scale down proportionally only when they exceed the page.
        factor = min(1.0, width_cm / sum(explicit_widths))
        widths = [value * factor for value in explicit_widths]
    table_width_cm = sum(widths)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    props = table._tbl.tblPr
    tw = props.find(qn('w:tblW'))
    tw.set(qn('w:type'), 'dxa')
    tw.set(qn('w:w'), str(int(Cm(table_width_cm).twips)))
    borders = element('tblBorders')
    for edge in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV'):
        borders.append(element(edge, val='single', sz=4, color=BORDER_COLOR))
    props.append(borders)
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for ridx, (row, values) in enumerate(zip(table.rows, rows)):
        if ridx == 0:
            row._tr.get_or_add_trPr().append(element('tblHeader', val='true'))
        # Intentionally allow row splitting; never set exact height or table-wide keepNext.
        for cidx, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            tcpr.append(element('shd', val='clear', fill=HEADER_FILL if ridx == 0 else (ROW_FILL if ridx % 2 else 'FFFFFF')))
            margins = element('tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                margins.append(element(edge, w=90 if edge in ('top', 'bottom') else 100, type='dxa'))
            tcpr.append(margins)
            p = cell.paragraphs[0]
            p.style = 'Guide Table'
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.keep_together = False
            if cidx == 0 or widths[cidx] < 2.4:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for lidx, text_line in enumerate(re.split(r'<br\s*/?>', text)):
                if lidx:
                    p.add_run().add_break()
                inline(p, text_line)
            if ridx == 0:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string('FFFFFF')
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.add_run().font.size = Pt(2)


def build(source, output, template=DEFAULT_TEMPLATE):
    output, template = Path(output), Path(template)
    if not template.is_file():
        raise FileNotFoundError(f'Template not found: {template}')
    blocks = parse_markdown(source)
    doc = Document(template)
    if any(kind == 'literal' for kind, _ in blocks) and 'Guide Code' not in doc.styles:
        style = doc.styles.add_style('Guide Code', WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles['Normal']
        style.font.name = 'Courier New'
        style.font.size = Pt(10)
        style.paragraph_format.left_indent = Cm(.25)
        style.paragraph_format.right_indent = Cm(.25)
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.element.get_or_add_pPr().append(element('shd', val='clear', fill='F2F4F7'))
    # The template has neutral demonstration slots; each is replaceable.
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'):
            doc.element.body.remove(child)
    title = next((str(value) for kind, value in blocks if kind == 'heading1'), '')
    doc.core_properties.title = title
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.comments = ''
    pending_widths = None
    for kind, value in blocks:
        if kind == 'widths':
            pending_widths = value
        elif kind == 'table':
            add_table(doc, value, pending_widths)
            pending_widths = None
        elif kind == 'pagebreak':
            doc.add_page_break()
        elif kind.startswith('heading'):
            style = {'heading1': 'Title', 'heading2': 'Heading 1', 'heading3': 'Heading 2'}[kind]
            p = doc.add_paragraph(style=style)
            inline(p, value)
        elif kind in ('prompt', 'literal'):
            # Keep content literal and let long blocks flow across pages.
            p = doc.add_paragraph(style='Guide Prompt' if kind == 'prompt' else 'Guide Code')
            p.add_run(value)
            p.paragraph_format.keep_together = False
            p.paragraph_format.keep_with_next = False
        else:
            style = {'paragraph': 'Normal', 'label': 'Normal', 'bullet': 'Guide Bullet', 'note': 'Guide Note'}[kind]
            p = doc.add_paragraph(style=style)
            if kind == 'label':
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.keep_together = True
            if kind == 'bullet':
                p.add_run('• ')
            inline(p, value)
    # Save to memory, then create exclusively: neither validation errors nor races overwrite files.
    buffer = BytesIO()
    doc.save(buffer)
    candidate, version = output, 1
    while True:
        try:
            file = candidate.open('xb')
        except FileExistsError:
            version += 1
            candidate = output.with_name(f'{output.stem}_v{version}{output.suffix}')
            continue
        with file:
            file.write(buffer.getvalue())
        return candidate


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path, metavar='INPUT.md')
    parser.add_argument('output', type=Path, metavar='OUTPUT.docx')
    parser.add_argument('--template', type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args()
    try:
        result = build(args.input.read_text(encoding='utf-8-sig'), args.output, args.template)
    except (OSError, ValueError) as exc:
        parser.exit(2, f'Error: {exc}\n')
    print(result)


if __name__ == '__main__':
    main()
