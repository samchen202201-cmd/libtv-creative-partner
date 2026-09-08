"""Focused regressions for Markdown fences and non-destructive output naming."""
from __future__ import annotations

import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document


SCRIPT = Path(__file__).resolve().parents[1] / 'scripts' / 'build_guide.py'
SPEC = importlib.util.spec_from_file_location('build_guide', SCRIPT)
builder = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(builder)


class BuildGuideTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_unknown_fences_preserve_literal_text_and_use_distinct_style(self):
        content = '  中文 **原样** | <br>\n\n\tprint("hello")  '
        for label in ('', 'text', 'markdown', 'python', 'custom-format'):
            with self.subTest(label=label):
                output = builder.build(f'# 指南\n\n```{label}\n{content}\n```', self.root / f'literal_{label or "empty"}.docx')
                paragraph = Document(output).paragraphs[-1]
                self.assertEqual(paragraph.text, content)
                self.assertEqual(paragraph.style.name, 'Guide Code')
                self.assertNotEqual(paragraph.style.name, 'Guide Prompt')

    def test_prompt_keeps_its_dedicated_style(self):
        output = builder.build('# 指南\n\n```prompt\n向右平移，保留音效。\n```', self.root / 'prompt.docx')
        paragraph = Document(output).paragraphs[-1]
        self.assertEqual(paragraph.style.name, 'Guide Prompt')
        self.assertEqual(paragraph.text, '向右平移，保留音效。')

    def test_long_fence_can_preserve_inner_triple_fence(self):
        content = '```prompt\n中文示例\n```'
        self.assertEqual(builder.parse_markdown(f'````markdown\n{content}\n````'), [('literal', content)])

    def test_malformed_input_reports_line_and_creates_no_output(self):
        samples = (
            ('# 指南\n\n```text\n未闭合', 'Unclosed code fence at line 3'),
            ('# 指南\n\n| A | B |\n| --- | --- |\n| one |', 'Table column count differs at line 5'),
        )
        for source, message in samples:
            with self.subTest(message=message):
                output = self.root / 'invalid.docx'
                with self.assertRaisesRegex(ValueError, message):
                    builder.build(source, output)
                self.assertFalse(output.exists())

    def test_same_name_gets_incremented_without_changing_earlier_files(self):
        output = self.root / '制作指南.docx'
        first = builder.build('# 原稿', output)
        original_hash = hashlib.sha256(first.read_bytes()).hexdigest()
        second = builder.build('# 返修一', output)
        second_hash = hashlib.sha256(second.read_bytes()).hexdigest()
        third = builder.build('# 返修二', output)
        self.assertEqual(first, output)
        self.assertEqual(second, self.root / '制作指南_v2.docx')
        self.assertEqual(third, self.root / '制作指南_v3.docx')
        self.assertEqual(hashlib.sha256(first.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(hashlib.sha256(second.read_bytes()).hexdigest(), second_hash)
        self.assertEqual(Document(third).paragraphs[0].text, '返修二')

    def test_exclusive_create_retries_when_another_writer_wins(self):
        output = self.root / 'race.docx'
        original_open = Path.open
        won = False

        def competing_open(path, mode='r', *args, **kwargs):
            nonlocal won
            if path == output and mode == 'xb' and not won:
                won = True
                with original_open(path, 'xb') as file:
                    file.write(b'other writer output')
            return original_open(path, mode, *args, **kwargs)

        with patch.object(Path, 'open', competing_open):
            result = builder.build('# 本次指南', output)
        self.assertTrue(won)
        self.assertEqual(result, self.root / 'race_v2.docx')
        self.assertEqual(output.read_bytes(), b'other writer output')
        self.assertEqual(Document(result).paragraphs[0].text, '本次指南')

    def test_other_write_errors_are_not_mistaken_for_name_collisions(self):
        output = self.root / 'denied.docx'
        original_open = Path.open

        def denied_open(path, mode='r', *args, **kwargs):
            if path.parent == self.root and mode == 'xb':
                raise PermissionError('write denied')
            return original_open(path, mode, *args, **kwargs)

        with patch.object(Path, 'open', denied_open):
            with self.assertRaisesRegex(PermissionError, 'write denied'):
                builder.build('# 指南', output)
        self.assertFalse(output.exists())

    def test_column_widths_fit_page_and_keep_narrative_first_column_useful(self):
        rows = [
            ['镜号', '画面与动作', '声音'],
            ['01', '人物从走廊尽头进入画面，停在门前并抬头观察门牌。', '脚步声、空调底噪'],
            ['02', '近景保持人物视线方向，门内暖光从右侧切入脸部。', '门轴轻响'],
        ]
        widths = builder.column_widths(rows, 24.0)
        self.assertEqual(len(widths), 3)
        self.assertTrue(all(width > 0 for width in widths))
        self.assertLessEqual(sum(widths), 24.0 + 1e-9)
        self.assertLessEqual(widths[0], 24.0 * 0.20 + 1e-9)

    def test_compact_table_does_not_get_equal_widths(self):
        rows = [['镜号', '时长', '景别'], ['01', '2秒', '近景'], ['02', '4秒', '全景']]
        widths = builder.column_widths(rows, 24.0)
        self.assertEqual(len(widths), 3)
        self.assertNotEqual(widths[0], widths[1])

    def test_explicit_width_comment_is_parsed_as_a_width_block(self):
        blocks = builder.parse_markdown(
            '# 指南\n\n<!-- widths: 4, 18, 5 -->\n\n| 镜号 | 画面 | 声音 |\n| --- | --- | --- |\n| 01 | 走廊 | 脚步 |'
        )
        self.assertEqual(blocks[1], ('widths', [4.0, 18.0, 5.0]))


if __name__ == '__main__':
    unittest.main()
